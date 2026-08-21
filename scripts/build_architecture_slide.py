"""One-slide architecture topology for the Shift-Left Provisioning Portal.

Built with python-pptx as NATIVE shapes (not a picture), so the diagram can be
edited in PowerPoint — boxes moved, wording changed — without regenerating it.

The diagram's job is to make the governing idea visible at a glance:
SEPARATION OF AUTHORITY. The portal collects and validates, Jira holds the
approval, and the orchestrator executes only after re-verifying that approval
itself. The approval gate is therefore the visual centre, not an afterthought.

No LibreOffice in this environment, so geometry is verified arithmetically by
scripts/check_slide_geometry.py rather than by eye.
"""

from __future__ import annotations

import pathlib

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Emu, Inches, Pt

OUT = pathlib.Path(__file__).resolve().parents[1] / "Infra-Portal-Architecture.pptx"

# Midnight Executive, with one warm accent reserved for the approval gate so the
# eye lands on the control point.
NAVY = RGBColor(0x1E, 0x27, 0x61)
ICE = RGBColor(0xCA, 0xDC, 0xFC)
ICE_DEEP = RGBColor(0x9D, 0xB8, 0xE8)
GOLD = RGBColor(0xE8, 0xA3, 0x3D)
GOLD_PALE = RGBColor(0xFB, 0xEC, 0xD2)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
SLATE = RGBColor(0x4A, 0x52, 0x6B)
MUTED = RGBColor(0x8A, 0x92, 0xA6)
PALE = RGBColor(0xF4, 0xF6, 0xFB)
GREEN = RGBColor(0x2C, 0x7A, 0x4B)

BODY = "Calibri"
HEAD = "Calibri"

prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)
slide = prs.slides.add_slide(prs.slide_layouts[6])   # blank

SHAPES: list[tuple[str, float, float, float, float]] = []   # for geometry checks
TEXTS: list[tuple[str, float, int, float, float]] = []      # name, pt, chars, width, height


def _record(name, l, t, w, h):
    SHAPES.append((name, l, t, w, h))


def textbox(name, l, t, w, h, runs, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP,
            space_after=2):
    """runs = [(text, size, bold, color), ...] — one paragraph each."""
    box = slide.shapes.add_textbox(Inches(l), Inches(t), Inches(w), Inches(h))
    tf = box.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = Emu(0)
    tf.margin_top = tf.margin_bottom = Emu(0)
    tf.vertical_anchor = anchor
    for i, (text, size, bold, color) in enumerate(runs):
        TEXTS.append((f"{name}[{i}]", size, len(text), w, h))
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        p.space_after = Pt(space_after)
        r = p.add_run()
        r.text = text
        r.font.size = Pt(size)
        r.font.bold = bold
        r.font.name = BODY
        r.font.color.rgb = color
    _record(name, l, t, w, h)
    return box


def card(name, l, t, w, h, fill, line, radius=0.06, shadow=False):
    sh = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE,
                                Inches(l), Inches(t), Inches(w), Inches(h))
    sh.adjustments[0] = radius
    sh.fill.solid()
    sh.fill.fore_color.rgb = fill
    sh.line.color.rgb = line
    sh.line.width = Pt(1.25)
    sh.shadow.inherit = shadow
    sh.text_frame.text = ""
    _record(name, l, t, w, h)
    return sh


def arrow(name, l, t, w, h, color=ICE_DEEP):
    sh = slide.shapes.add_shape(MSO_SHAPE.RIGHT_ARROW,
                                Inches(l), Inches(t), Inches(w), Inches(h))
    sh.fill.solid()
    sh.fill.fore_color.rgb = color
    sh.line.fill.background()
    sh.shadow.inherit = False
    _record(name, l, t, w, h)
    return sh


# ---------------------------------------------------------------- background --
bg = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0,
                            prs.slide_width, prs.slide_height)
bg.fill.solid()
bg.fill.fore_color.rgb = WHITE
bg.line.fill.background()
bg.shadow.inherit = False

# --------------------------------------------------------------------- title --
textbox("title", 0.55, 0.34, 12.2, 0.55, [
    ("Shift-Left Infrastructure Provisioning Portal — Architecture", 28, True, NAVY),
])
textbox("subtitle", 0.55, 0.92, 12.2, 0.34, [
    ("Self-service request  →  governed approval  →  automated build, "
     "with every privileged action audited", 13, False, SLATE),
])

# ------------------------------------------------------- the five zones (L→R) --
COL_TOP = 1.72
COL_H = 3.62
GAP_ARROW_W = 0.30

ZONES = [
    # (key, x, width, header, fill, line, items)
    ("users", 0.55, 1.95, "CONSUMERS", PALE, ICE_DEEP, [
        ("Requester", "raises and tracks requests"),
        ("Approver", "line manager / platform"),
        ("Platform admin", "catalogue, settings, RBAC"),
    ]),
    ("portal", 2.95, 3.05, "PORTAL  ·  API  ·  POLICY", ICE, NAVY, [
        ("React + IBM Carbon", "guided form, drafts, My Environments"),
        ("FastAPI API", "validation, sizing, pricing — server-side"),
        ("Open Policy Agent", "policy evaluated on every request"),
        ("PostgreSQL 16", "state + append-only audit chain"),
    ]),
    ("approval", 6.45, 2.30, "APPROVAL AUTHORITY", GOLD_PALE, GOLD, [
        ("Jira Service Management", "live — project SDIMD"),
        ("Holds the approval", "the portal never self-approves"),
    ]),
    ("orch", 9.20, 2.30, "ORCHESTRATOR", ICE, NAVY, [
        ("Leader-elected poller", "single active replica"),
        ("Terraform runner", "per-request isolated state"),
        ("Re-verifies approval", "before it builds anything"),
    ]),
    ("clouds", 11.95, 0.83, "TARGETS", PALE, ICE_DEEP, []),
]

for key, x, w, header, fill, line, items in ZONES:
    card(f"zone_{key}", x, COL_TOP, w, COL_H, fill, line,
         shadow=(key == "approval"))
    textbox(f"hdr_{key}", x + 0.16, COL_TOP + 0.16, w - 0.32, 0.28,
            [(header, 10.5, True, NAVY if key != "approval" else
              RGBColor(0x8A, 0x5A, 0x10))],
            align=PP_ALIGN.CENTER)
    y = COL_TOP + 0.60
    for name, detail in items:
        textbox(f"itm_{key}_{name[:10]}", x + 0.16, y, w - 0.32, 0.62, [
            (name, 11.5, True, NAVY),
            (detail, 9.5, False, SLATE),
        ], space_after=0)
        y += 0.74

# targets column: stacked chips (kept narrow, so text is short by design)
TARGETS = [
    ("OCI", "LIVE", GREEN, WHITE),
    ("Azure", "NEXT", GOLD, RGBColor(0x5A, 0x3B, 0x08)),
    ("AWS", "cost only", None, MUTED),
    ("GCP", "cost only", None, MUTED),
    ("On-prem", "cost only", None, MUTED),
]
ty = COL_TOP + 0.60
for name, tag, chip, textcol in TARGETS:
    if chip is not None:
        c = card(f"chip_{name}", 12.06, ty, 0.61, 0.50, chip, chip, radius=0.15)
        textbox(f"chiptxt_{name}", 12.06, ty + 0.05, 0.61, 0.20,
                [(name, 10.5, True, textcol)], align=PP_ALIGN.CENTER)
        textbox(f"chiptag_{name}", 12.06, ty + 0.26, 0.61, 0.18,
                [(tag, 7.5, True, textcol)], align=PP_ALIGN.CENTER)
    else:
        textbox(f"chiptxt_{name}", 12.06, ty + 0.04, 0.61, 0.20,
                [(name, 10.5, True, MUTED)], align=PP_ALIGN.CENTER)
        textbox(f"chiptag_{name}", 12.06, ty + 0.24, 0.61, 0.18,
                [(tag, 7.5, False, MUTED)], align=PP_ALIGN.CENTER)
    ty += 0.60

# ------------------------------------------------------------------- arrows --
ARROW_Y = COL_TOP + 1.55
for ax in (2.56, 6.06, 8.81, 11.56):
    arrow(f"arrow_{ax}", ax, ARROW_Y, GAP_ARROW_W, 0.42)

# ------------------------------------------------- the governing idea, called out
card("authority_band", 0.55, 5.62, 8.35, 0.72, NAVY, NAVY, radius=0.10)
textbox("authority_text", 0.78, 5.76, 7.90, 0.46, [
    ("SEPARATION OF AUTHORITY", 11, True, GOLD),
    ("The portal collects and validates  ·  Jira holds the approval  ·  "
     "the orchestrator executes only after re-verifying it", 10.5, False, WHITE),
], space_after=1)

# ------------------------------------------------------- cross-cutting services
card("ai_card", 9.20, 5.62, 3.58, 0.72, PALE, ICE_DEEP, radius=0.10)
textbox("ai_text", 9.34, 5.72, 3.30, 0.58, [
    ("AI ADVISOR", 11, True, NAVY),
    ("Recommends only — never approves, never provisions, holds no credentials",
     9.5, False, SLATE),
], space_after=1)

# --------------------------------------------------------------------- footer --
textbox("footer", 0.55, 6.62, 12.2, 0.30, [
    ("Identity: Microsoft Entra ID (sign-in)   ·   Secrets: OCI Vault — the portal "
     "holds a reference, never a password   ·   Audit: append-only, hash-chained   "
     "·   Mock mode is the default; nothing real builds without an approval",
     8.5, False, MUTED),
])
textbox("asof", 0.55, 7.00, 12.2, 0.26, [
    ("As built, 21 Aug 2026  ·  9 OCI workload types provisioning end to end",
     8.5, True, MUTED),
])

prs.save(OUT)
print(f"written: {OUT}")

# Hand the geometry to the checker.
import json
(pathlib.Path(__file__).resolve().parent / "_slide_geometry.json").write_text(
    json.dumps({"slide": [13.333, 7.5], "shapes": SHAPES,
                "texts": TEXTS}, indent=1), encoding="utf-8")
print(f"recorded {len(SHAPES)} shapes for geometry checking")
